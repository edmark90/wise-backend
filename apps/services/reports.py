"""Report generation service (Excel + PDF).

Builds monthly/yearly report bundles for the admin Reports page:
  - Waste Classification (waste_records joined with users)
  - Collection Schedule (collection_schedule)
  - Users (users)

Excel output is a single .xlsx with one sheet per selected report type
(prefixed with a Summary sheet). PDF output is a single branded summary
page (Muntinlupa header + stat cards).
"""
import calendar
import io
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from sqlalchemy import func
from sqlalchemy.orm import Session

from apps.models.collection_schedule import CollectionSchedule
from apps.services.collection_schedule import derive_schedule_status
from apps.models.user import User
from apps.models.waste_record import WasteRecord
from apps.utils.ph_time import ph_now

REPORT_TYPES = ["waste", "schedule", "users"]

TYPE_LABELS = {
    "waste": "Waste Classification",
    "schedule": "Collection Schedule",
    "users": "Users",
}

TYPE_SHEET_NAMES = {
    "waste": "Waste Records",
    "schedule": "Collection Schedule",
    "users": "Users",
}

LOW_CONFIDENCE_THRESHOLD = 50.0

HEADER_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(size=10.5)
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def resolve_period(year: int, month: int | None) -> tuple[datetime, datetime]:
    """Return (start, end) naive datetimes for the requested period.

    month=None means the whole year (Jan 1 .. Dec 31).
    """
    if month:
        last_day = calendar.monthrange(year, month)[1]
        start = datetime(year, month, 1, 0, 0, 0)
        end = datetime(year, month, last_day, 23, 59, 59)
    else:
        start = datetime(year, 1, 1, 0, 0, 0)
        end = datetime(year, 12, 31, 23, 59, 59)
    return start, end


def period_label(year: int, month: int | None) -> str:
    if month:
        return datetime(year, month, 1).strftime("%B %Y")
    return str(year)


# ---------------------------------------------------------------------------
# Data gathering
# ---------------------------------------------------------------------------

def _waste_rows(db: Session, start: datetime, end: datetime):
    rows = (
        db.query(
            WasteRecord.classified_at,
            User.fullname,
            User.barangay,
            User.zone,
            WasteRecord.waste_type,
            WasteRecord.disposal_category,
            WasteRecord.confidence,
            WasteRecord.is_flagged,
        )
        .join(User, WasteRecord.user_id == User.id)
        .filter(WasteRecord.classified_at >= start, WasteRecord.classified_at <= end)
        .order_by(WasteRecord.classified_at.asc())
        .all()
    )
    return [
        [
            r.classified_at.strftime("%Y-%m-%d %H:%M") if r.classified_at else "-",
            r.fullname or "-",
            f"{r.barangay or ''}{', ' + r.zone if r.zone else ''}".strip(" ,"),
            r.waste_type,
            r.disposal_category,
            f"{float(r.confidence):.1f}%" if r.confidence is not None else "-",
            "Flagged" if r.is_flagged else "OK",
        ]
        for r in rows
    ]


def _schedule_rows(db: Session, start: datetime, end: datetime):
    rows = (
        db.query(CollectionSchedule)
        .filter(CollectionSchedule.collection_date >= start.date(), CollectionSchedule.collection_date <= end.date())
        .order_by(CollectionSchedule.collection_date.asc(), CollectionSchedule.collection_time.asc())
        .all()
    )
    return [
        [
            r.collection_date.strftime("%Y-%m-%d") if r.collection_date else "-",
            r.collection_time.strftime("%I:%M %p") if r.collection_time else "-",
            r.barangay,
            r.zone or "",
            r.route_name or "-",
            r.starting_point or "-",
            r.assigned_personnel or "-",
            derive_schedule_status(r),
            r.remarks or "",
        ]
        for r in rows
    ]


def _user_rows(db: Session, start: datetime, end: datetime):
    rows = (
        db.query(User)
        .filter(User.created_at >= start, User.created_at <= end)
        .order_by(User.created_at.asc())
        .all()
    )
    return [
        [
            r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "-",
            r.fullname,
            r.email,
            r.phone or "-",
            r.role.capitalize(),
            r.barangay or "",
            r.zone or "",
            "Active" if r.is_active else "Inactive",
        ]
        for r in rows
    ]


def _notification_rows(db: Session, start: datetime, end: datetime):
    return []


def _data_for(db: Session, start: datetime, end: datetime, types: list[str]):
    data = {}
    for t in types:
        if t == "waste":
            data[t] = _waste_rows(db, start, end)
        elif t == "schedule":
            data[t] = _schedule_rows(db, start, end)
        elif t == "users":
            data[t] = _user_rows(db, start, end)
        elif t == "notifications":
            data[t] = []
    return data


# ---------------------------------------------------------------------------
# Aggregate stats
# ---------------------------------------------------------------------------

def waste_stats(db: Session, start: datetime, end: datetime) -> dict:
    """Summary analytics for waste records in the period."""
    rows = (
        db.query(
            WasteRecord.waste_type,
            func.count(WasteRecord.id).label("cnt"),
            func.avg(WasteRecord.confidence).label("avg_conf"),
        )
        .filter(WasteRecord.classified_at >= start, WasteRecord.classified_at <= end)
        .group_by(WasteRecord.waste_type)
        .all()
    )

    total = sum(int(r.cnt) for r in rows)
    total_records = (
        db.query(func.count(WasteRecord.id))
        .filter(WasteRecord.classified_at >= start, WasteRecord.classified_at <= end)
        .scalar() or 0
    )
    avg_conf = None
    conf_counts = []
    for r in rows:
        if r.avg_conf is not None:
            conf_counts.append(float(r.avg_conf) * int(r.cnt))
    if conf_counts:
        avg_conf = round(sum(conf_counts) / max(sum(int(r.cnt) for r in rows), 1), 1)

    low_conf = (
        db.query(func.count(WasteRecord.id))
        .filter(
            WasteRecord.classified_at >= start,
            WasteRecord.classified_at <= end,
            WasteRecord.confidence.isnot(None),
            WasteRecord.confidence < LOW_CONFIDENCE_THRESHOLD,
        )
        .scalar() or 0
    )

    per_class = {}
    most_common_type = None
    most_common_count = 0
    for r in rows:
        per_class[r.waste_type] = int(r.cnt)
        if int(r.cnt) > most_common_count:
            most_common_count = int(r.cnt)
            most_common_type = r.waste_type

    return {
        "total_records": total_records,
        "total_classified": total,
        "avg_confidence": avg_conf,
        "low_confidence": low_conf,
        "per_class": per_class,
        "most_common_type": most_common_type,
        "most_common_count": most_common_count,
    }


def schedule_stats(db: Session, start: datetime, end: datetime) -> dict:
    """Route completion analytics for the period (server-time-derived statuses)."""
    rows = (
        db.query(CollectionSchedule)
        .filter(CollectionSchedule.collection_date >= start.date(), CollectionSchedule.collection_date <= end.date())
        .all()
    )

    counts = {}
    for r in rows:
        status = derive_schedule_status(r)
        counts[status] = counts.get(status, 0) + 1

    total = sum(counts.values())
    completed = counts.get("Completed", 0)
    cancelled = counts.get("Cancelled", 0)
    delayed = counts.get("Delayed", 0)
    upcoming = counts.get("Upcoming", 0)
    completion_rate = round(completed / total * 100, 1) if total else 0.0

    return {
        "total_routes": total,
        "completed": completed,
        "cancelled": cancelled,
        "delayed": delayed,
        "upcoming": upcoming,
        "completion_rate": completion_rate,
        "status_breakdown": counts,
    }


# ---------------------------------------------------------------------------
# Excel (.xlsx)
# ---------------------------------------------------------------------------

EXCEL_HEADERS = {
    "waste": ["Date & Time", "Citizen", "Barangay / Zone", "Waste Type", "Classification", "Confidence", "Status"],
    "schedule": ["Date", "Time", "Barangay", "Zone", "Route", "Starting Point", "Personnel", "Status", "Remarks"],
    "users": ["Registered On", "Full Name", "Email", "Phone", "Role", "Barangay", "Zone", "Status"],
}


def build_excel(
    db: Session,
    year: int,
    month: int | None,
    types: list[str],
    generated_at: datetime | None = None,
) -> bytes:
    """Generate a single .xlsx workbook with a Summary sheet + one sheet per type."""
    generated_at = generated_at or ph_now()
    start, end = resolve_period(year, month)
    label = period_label(year, month)
    data = _data_for(db, start, end, types)
    waste = waste_stats(db, start, end)
    schedule = schedule_stats(db, start, end)

    wb = Workbook()

    # ---- Summary sheet -------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False

    title_cell = ws.cell(row=1, column=1, value="WISE SYSTEM — Report")
    title_cell.font = Font(bold=True, size=18, color="1E293B")
    period_cell = ws.cell(row=2, column=1, value=f"Period: {label}")
    period_cell.font = Font(size=13, color="64748B")

    row = 4
    sections = []

    if "waste" in types:
        w_rows = []
        w_rows.append(("Total Waste Records", waste["total_records"]))
        w_rows.append(("Average AI Confidence", f"{waste['avg_confidence']}%" if waste["avg_confidence"] is not None else "-"))
        w_rows.append((f"Low Accuracy (<{int(LOW_CONFIDENCE_THRESHOLD)}%)", waste["low_confidence"]))
        w_rows.append(("Most Common Waste Type", waste["most_common_type"] or "-"))
        for cls, cnt in waste["per_class"].items():
            w_rows.append((f"  • {cls}", cnt))
        sections.append(("Waste Classification", w_rows))

    if "schedule" in types:
        s_rows = [
            ("Total Routes Scheduled", schedule["total_routes"]),
            ("Routes Completed", schedule["completed"]),
            ("Routes Cancelled", schedule["cancelled"]),
            ("Routes Delayed", schedule["delayed"]),
            ("Routes Upcoming", schedule["upcoming"]),
            ("Completion Rate", f"{schedule['completion_rate']}%"),
        ]
        sections.append(("Collection Schedule", s_rows))

    if "users" in types:
        u_rows = [("New Users Registered", len(data.get("users", [])))]
        sections.append(("Users", u_rows))

    for section_name, rows in sections:
        ws.cell(row=row, column=1, value=section_name).font = Font(bold=True, size=12, color="2563EB")
        row += 1
        for key, val in rows:
            ws.cell(row=row, column=1, value=key).font = BODY_FONT
            ws.cell(row=row, column=2, value=val).font = BODY_FONT
            ws.cell(row=row, column=1).border = BORDER
            ws.cell(row=row, column=2).border = BORDER
            row += 1
        row += 1

    gen_row = row + 1
    ws.cell(row=gen_row, column=1, value=f"Generated: {generated_at.strftime('%Y-%m-%d %H:%M')} (Manila)").font = Font(
        size=10, color="94A3B8"
    )

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16

    # ---- Per-type sheets ------------------------------------------------
    for t in types:
        headers = EXCEL_HEADERS[t]
        rows = data.get(t, [])
        ws = wb.create_sheet(TYPE_SHEET_NAMES[t])

        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center")

        for r, row in enumerate(rows, start=2):
            for c, val in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c, value=val)
                cell.font = BODY_FONT
                cell.border = BORDER
                cell.alignment = Alignment(vertical="top")

        widths = [22, 22, 18, 16, 16, 12, 12, 14, 18]
        for c in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(c)].width = widths[c - 1] if c - 1 < len(widths) else 14

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 2)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

_NAVY = colors.HexColor("#0F1B2D")
_BLUE = colors.HexColor("#2563EB")
_MUTED = colors.HexColor("#64748B")

_TITLE_STYLE = ParagraphStyle(
    "Title", fontName="Helvetica-Bold", fontSize=21, leading=25, textColor=_NAVY,
    alignment=TA_CENTER, spaceAfter=2,
)
_PERIOD_STYLE = ParagraphStyle(
    "Period", fontName="Helvetica-Bold", fontSize=12, leading=16, textColor=_BLUE,
    alignment=TA_CENTER, spaceAfter=14,
)
_SECTION_STYLE = ParagraphStyle(
    "Section", fontName="Helvetica-Bold", fontSize=12.5, leading=15, textColor=_NAVY,
    spaceBefore=2, spaceAfter=2,
)
_CARD_VALUE_STYLE = ParagraphStyle(
    "CardValue", fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=_NAVY,
    alignment=TA_CENTER,
)
_CARD_LABEL_STYLE = ParagraphStyle(
    "CardLabel", fontName="Helvetica", fontSize=7.5, leading=9.5, textColor=_MUTED,
    alignment=TA_CENTER,
)
_GEN_STYLE = ParagraphStyle(
    "Gen", fontName="Helvetica", fontSize=8, leading=11, textColor=colors.HexColor("#94A3B8"),
    alignment=TA_CENTER,
)

_LOGO_PATH = os.path.join(os.path.dirname(__file__), "..", "assets", "wise-logo.png")


def _stat_card(value: str, label: str, color: str) -> Table:
    card = Table(
        [[Paragraph(value, _CARD_VALUE_STYLE)], [Paragraph(label, _CARD_LABEL_STYLE)]],
        colWidths=[36 * mm],
        rowHeights=[8 * mm, 5.5 * mm],
        style=[
            ("BACKGROUND", (0, 0), (-1, -1), color),
            ("ROUNDEDCORNERS", [2.5, 2.5, 2.5, 2.5]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ],
    )
    return card


def _section_heading(text: str) -> Table:
    head = Table(
        [[Paragraph(text, _SECTION_STYLE), Paragraph("", _SECTION_STYLE)]],
        colWidths=[70 * mm, 30 * mm],
        style=[
            ("BACKGROUND", (1, 0), (1, 0), _BLUE),
            ("ROUNDEDCORNERS", [0, 1.2, 1.2, 0]),
        ],
    )
    return head


def _pdf_background(canvas, doc):
    """Draw the deep-navy header band + footer on every page."""
    page_w = landscape(A4)[0]
    page_h = landscape(A4)[1]

    # Header band
    canvas.saveState()
    band_h = 30 * mm
    p = canvas.beginPath()
    p.moveTo(0, page_h)
    p.lineTo(page_w, page_h)
    p.lineTo(page_w, page_h - band_h)
    p.lineTo(0, page_h - band_h)
    p.close()
    canvas.setFillColor(_NAVY)
    canvas.drawPath(p, stroke=0, fill=1)

    # Accent underline under the band
    canvas.setFillColor(_BLUE)
    canvas.rect(0, page_h - band_h - 1.5 * mm, page_w, 1.5 * mm, stroke=0, fill=1)

    # Logo on the left of the band
    logo = _LOGO_PATH
    if os.path.exists(logo):
        try:
            from PIL import Image as PILImage
            im = PILImage.open(logo)
            iw, ih = im.size
            target_h = band_h - 8 * mm
            target_w = target_h * (iw / ih)
            canvas.drawImage(logo, 12 * mm, page_h - band_h + (band_h - target_h) / 2, width=target_w, height=target_h, mask="auto", preserveAspectRatio=True)
        except Exception:
            pass

    # Band text (right side)
    canvas.setFillColor(colors.white)
    canvas.setFont("Helvetica-Bold", 17)
    canvas.drawRightString(page_w - 14 * mm, page_h - 9 * mm, "MUNTINLUPA")
    canvas.setFont("Helvetica", 9)
    canvas.setFillColor(colors.HexColor("#93C5FD"))
    canvas.drawRightString(page_w - 14 * mm, page_h - 15 * mm, "Waste Information System for the Environment (WISE)")

    # Footer
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#94A3B8"))
    canvas.drawCentredString(page_w / 2, 7 * mm, f"WISE SYSTEM — Monthly Report")
    canvas.drawRightString(page_w - 12 * mm, 7 * mm, f"Page {doc.page}")
    canvas.restoreState()


def build_pdf(
    db: Session,
    year: int,
    month: int | None,
    types: list[str],
    generated_at: datetime | None = None,
) -> bytes:
    """Generate a single-page branded summary PDF for the chosen period.

    The output is a clean summary (no per-type detail tables): Muntinlupa
    header band, period, and stat cards for the selected report types.
    """
    generated_at = generated_at or ph_now()
    start, end = resolve_period(year, month)
    label = period_label(year, month)
    waste = waste_stats(db, start, end)
    schedule = schedule_stats(db, start, end)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=36 * mm,
        bottomMargin=14 * mm,
        title=f"WISE Report {label}",
        author="WISE System",
    )

    story = []

    story.append(Paragraph("Waste Management Report", _TITLE_STYLE))
    story.append(Paragraph(f"Reporting Period: {label}", _PERIOD_STYLE))

    # Waste Classification cards
    if "waste" in types:
        story.append(_section_heading("Waste Classification"))
        waste_cards = [
            _stat_card(str(waste["total_records"]), "Total Records", "#DBEAFE"),
            _stat_card(
                f"{waste['avg_confidence']}%" if waste["avg_confidence"] is not None else "-",
                "Average Accuracy",
                "#D1FAE5",
            ),
            _stat_card(str(waste["low_confidence"]), "Low Accuracy (<50%)", "#FEE2E2"),
            _stat_card(waste["most_common_type"] or "-", "Most Common Type", "#EDE9FE"),
        ]
        story.append(Table([waste_cards], colWidths=[36 * mm] * 4, rowHeights=[14 * mm]))
        story.append(Spacer(1, 8))

    # Collection Schedule cards
    if "schedule" in types:
        story.append(_section_heading("Collection Schedule"))
        schedule_cards = [
            _stat_card(str(schedule["total_routes"]), "Total Routes", "#DBEAFE"),
            _stat_card(str(schedule["completed"]), "Completed", "#D1FAE5"),
            _stat_card(str(schedule["cancelled"]), "Cancelled", "#FEE2E2"),
            _stat_card(str(schedule["delayed"]), "Delayed", "#FEF3C7"),
        ]
        story.append(Table([schedule_cards], colWidths=[36 * mm] * 4, rowHeights=[14 * mm]))
        rate_cards = [
            _stat_card(str(schedule["upcoming"]), "Upcoming", "#E0E7FF"),
            _stat_card(f"{schedule['completion_rate']}%", "Completion Rate", "#10B981"),
            _stat_card("", "", "#FFFFFF"),
            _stat_card("", "", "#FFFFFF"),
        ]
        story.append(Table([rate_cards], colWidths=[36 * mm] * 4, rowHeights=[14 * mm]))
        story.append(Spacer(1, 8))

    # Users
    if "users" in types:
        story.append(_section_heading("Users"))
        from apps.models.user import User as _UserModel
        users_total = (
            db.query(func.count(_UserModel.id))
            .filter(_UserModel.created_at >= start, _UserModel.created_at <= end)
            .scalar() or 0
        )
        user_cards = [
            _stat_card(str(users_total), "New Users Registered", "#DBEAFE"),
            _stat_card("", "", "#FFFFFF"),
            _stat_card("", "", "#FFFFFF"),
            _stat_card("", "", "#FFFFFF"),
        ]
        story.append(Table([user_cards], colWidths=[36 * mm] * 4, rowHeights=[14 * mm]))
        story.append(Spacer(1, 8))

    story.append(Spacer(1, 10))
    story.append(Paragraph(f"Generated: {generated_at.strftime('%Y-%m-%d %H:%M')} (Manila)", _GEN_STYLE))

    doc.build(story, onFirstPage=_pdf_background, onLaterPages=_pdf_background)
    return buf.getvalue()